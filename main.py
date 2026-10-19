import io, random, uuid
from pathlib import Path
from fastapi import FastAPI, Form, File, UploadFile, HTTPException
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.exc import SQLAlchemyError
from database import Base, CauHoi, SessionLocal, engine
from export_word import tao_file_word_de_thi

app=FastAPI(title="Ngân hàng câu hỏi và tạo đề thi")
Base.metadata.create_all(bind=engine)
BASE=Path(__file__).resolve().parent
VALID_TYPES={"Dang_I","Dang_II","Dang_III"}
VALID_LEVELS={"biet","hieu","vandung"}
TYPE_MAP={"dang_i":"Dang_I","dạng i":"Dang_I","trac_nghiem":"Dang_I","dang_ii":"Dang_II","dạng ii":"Dang_II","dung_sai":"Dang_II","đúng/sai":"Dang_II","dang_iii":"Dang_III","dạng iii":"Dang_III","tu_luan":"Dang_III","tự luận":"Dang_III"}
LEVEL_MAP={"biết":"biet","nhận biết":"biet","nhan_biet":"biet","hiểu":"hieu","thông hiểu":"hieu","thong_hieu":"hieu","vận dụng":"vandung","van_dung":"vandung"}
def norm_type(x):
    s=str(x or "").strip()
    return TYPE_MAP.get(s.lower(),s if s in VALID_TYPES else "")
def norm_level(x):
    s=str(x or "").strip()
    return LEVEL_MAP.get(s.lower(),s if s in VALID_LEVELS else "")
def txt(x): return "" if x is None else str(x).strip()
def as_dict(q):
    keys=["id","mon_hoc","khoi_lop","dang_cau_hoi","muc_do","noi_dung","dap_an_a","dap_an_b","dap_an_c","dap_an_d","dap_an_dung","giai_thich"]
    return {k:getattr(q,k) for k in keys}

@app.get("/")
def home(): return FileResponse(BASE/"index.html")
@app.get("/health")
def health(): return {"status":"ok"}

@app.get("/api/cau-hoi")
def list_questions(mon_hoc:str="",khoi_lop:str="",dang_cau_hoi:str="",muc_do:str=""):
    db=SessionLocal()
    try:
        q=db.query(CauHoi)
        if mon_hoc.strip(): q=q.filter(CauHoi.mon_hoc==mon_hoc.strip())
        if khoi_lop.strip(): q=q.filter(CauHoi.khoi_lop==khoi_lop.strip())
        if dang_cau_hoi: q=q.filter(CauHoi.dang_cau_hoi==norm_type(dang_cau_hoi))
        if muc_do: q=q.filter(CauHoi.muc_do==norm_level(muc_do))
        rows=q.order_by(CauHoi.id.desc()).limit(1000).all()
        return {"total":len(rows),"questions":[as_dict(x) for x in rows]}
    finally: db.close()

@app.post("/api/cau-hoi")
def add_question(mon_hoc:str=Form(...),khoi_lop:str=Form(...),dang_cau_hoi:str=Form(...),muc_do:str=Form(...),noi_dung:str=Form(...),
                 dap_an_a:str=Form(""),dap_an_b:str=Form(""),dap_an_c:str=Form(""),dap_an_d:str=Form(""),dap_an_dung:str=Form(""),giai_thich:str=Form("")):
    kind,level=norm_type(dang_cau_hoi),norm_level(muc_do)
    a,b,c,d=map(txt,[dap_an_a,dap_an_b,dap_an_c,dap_an_d])
    ans=txt(dap_an_dung)
    if not txt(mon_hoc) or not txt(khoi_lop) or not txt(noi_dung): raise HTTPException(400,"Môn, khối và nội dung câu hỏi là bắt buộc.")
    if kind not in VALID_TYPES or level not in VALID_LEVELS: raise HTTPException(400,"Dạng câu hỏi hoặc mức độ không hợp lệ.")
    if kind=="Dang_I" and (not all([a,b,c,d]) or ans.upper() not in {"A","B","C","D"}): raise HTTPException(400,"Dạng I cần đủ 4 lựa chọn và đáp án A/B/C/D.")
    if kind=="Dang_II" and (not all([a,b,c,d]) or not ans): raise HTTPException(400,"Dạng II cần đủ 4 mệnh đề và đáp án đúng.")
    db=SessionLocal()
    try:
        q=CauHoi(mon_hoc=txt(mon_hoc),khoi_lop=txt(khoi_lop),dang_cau_hoi=kind,muc_do=level,noi_dung=txt(noi_dung),dap_an_a=a,dap_an_b=b,dap_an_c=c,dap_an_d=d,dap_an_dung=ans,giai_thich=txt(giai_thich))
        db.add(q); db.commit(); db.refresh(q)
        return {"message":"Đã lưu câu hỏi.","question":as_dict(q)}
    except SQLAlchemyError as e:
        db.rollback(); raise HTTPException(500,f"Không lưu được: {e}")
    finally: db.close()

@app.delete("/api/cau-hoi/{qid}")
def delete_question(qid:int):
    db=SessionLocal()
    try:
        q=db.query(CauHoi).filter(CauHoi.id==qid).first()
        if not q: raise HTTPException(404,"Không tìm thấy câu hỏi.")
        db.delete(q); db.commit(); return {"message":"Đã xóa."}
    finally: db.close()

@app.get("/api/mau-excel")
def excel_template():
    wb=Workbook(); ws=wb.active; ws.title="NganHangCauHoi"
    ws.append(["mon_hoc","khoi_lop","dang_cau_hoi","muc_do","noi_dung","dap_an_a","dap_an_b","dap_an_c","dap_an_d","dap_an_dung","giai_thich"])
    ws.append(["Tin học","10","Dang_I","biet","CPU có chức năng gì?","Xử lý dữ liệu","In văn bản","Nhập dữ liệu","Lưu điện","A",""])
    ws.append(["Tin học","10","Dang_II","hieu","Xác định đúng/sai cho các mệnh đề sau.","CPU là bộ xử lý","RAM là bộ nhớ tạm","Bàn phím là thiết bị xuất","Màn hình là thiết bị nhập","Đúng,Đúng,Sai,Sai",""])
    ws.append(["Tin học","10","Dang_III","vandung","Viết chương trình giải bài toán...","Đáp án mẫu/hướng dẫn chấm","","","","",""])
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=24
    ws.column_dimensions["E"].width=50; ws.column_dimensions["K"].width=40
    stream=io.BytesIO(); wb.save(stream); stream.seek(0)
    return StreamingResponse(stream,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":'attachment; filename="Mau_Ngan_Hang_Cau_Hoi.xlsx"'})

@app.post("/api/import-excel")
async def import_excel(file:UploadFile=File(...)):
    if not file.filename or not file.filename.lower().endswith((".xlsx",".xlsm")): raise HTTPException(400,"Chỉ hỗ trợ Excel .xlsx/.xlsm.")
    try:
        wb=load_workbook(io.BytesIO(await file.read()),read_only=True,data_only=True)
        ws=wb["NganHangCauHoi"] if "NganHangCauHoi" in wb.sheetnames else wb.active
        rows=ws.iter_rows(values_only=True); headers=[txt(v).lower() for v in next(rows,[])]
        required=["mon_hoc","khoi_lop","dang_cau_hoi","muc_do","noi_dung"]
        if any(k not in headers for k in required): raise HTTPException(400,"Thiếu cột bắt buộc: "+", ".join(k for k in required if k not in headers))
        idx={h:i for i,h in enumerate(headers)}; db=SessionLocal(); added=skipped=0; errors=[]
        try:
            for rn,row in enumerate(rows,start=2):
                if not row or not any(txt(v) for v in row): continue
                data={h:(row[i] if i<len(row) else None) for h,i in idx.items()}
                kind,level=norm_type(data.get("dang_cau_hoi")),norm_level(data.get("muc_do"))
                mon,grade,content=txt(data.get("mon_hoc")),txt(data.get("khoi_lop")),txt(data.get("noi_dung"))
                a,b,c,d=[txt(data.get(k)) for k in ["dap_an_a","dap_an_b","dap_an_c","dap_an_d"]]
                ans,expl=txt(data.get("dap_an_dung")),txt(data.get("giai_thich"))
                error=""
                if not mon or not grade or not content or kind not in VALID_TYPES or level not in VALID_LEVELS: error="Thiếu dữ liệu hoặc dạng/mức độ không hợp lệ."
                elif kind=="Dang_I" and (not all([a,b,c,d]) or ans.upper() not in {"A","B","C","D"}): error="Dạng I cần đủ A/B/C/D và đáp án đúng A/B/C/D."
                elif kind=="Dang_II" and (not all([a,b,c,d]) or not ans): error="Dạng II cần đủ 4 mệnh đề và đáp án đúng."
                if error: skipped+=1; errors.append(f"Dòng {rn}: {error}"); continue
                db.add(CauHoi(mon_hoc=mon,khoi_lop=grade,dang_cau_hoi=kind,muc_do=level,noi_dung=content,dap_an_a=a,dap_an_b=b,dap_an_c=c,dap_an_d=d,dap_an_dung=ans,giai_thich=expl)); added+=1
            db.commit()
        except Exception: db.rollback(); raise
        finally: db.close()
        return {"message":f"Đã nhập {added} câu; bỏ qua {skipped} dòng.","added":added,"skipped":skipped,"errors":errors[:50]}
    except HTTPException: raise
    except Exception as e: raise HTTPException(400,f"Không đọc được Excel: {e}")

@app.post("/api/tao-de")
def create_exam(mon_hoc:str=Form(...),khoi_lop:str=Form(...),loai_kiem_tra:str=Form(...),thoi_gian:int=Form(45),
                dang1_biet:int=Form(0),dang1_hieu:int=Form(0),dang1_vandung:int=Form(0),
                dang2_biet:int=Form(0),dang2_hieu:int=Form(0),dang2_vandung:int=Form(0),
                dang3_biet:int=Form(0),dang3_hieu:int=Form(0),dang3_vandung:int=Form(0),kem_dap_an:str=Form("yes")):
    matrix={"Dang_I":{"biet":dang1_biet,"hieu":dang1_hieu,"vandung":dang1_vandung},"Dang_II":{"biet":dang2_biet,"hieu":dang2_hieu,"vandung":dang2_vandung},"Dang_III":{"biet":dang3_biet,"hieu":dang3_hieu,"vandung":dang3_vandung}}
    if thoi_gian<1: raise HTTPException(400,"Thời gian phải lớn hơn 0.")
    if any(n<0 for group in matrix.values() for n in group.values()): raise HTTPException(400,"Số lượng câu không được âm.")
    if not sum(n for group in matrix.values() for n in group.values()): raise HTTPException(400,"Hãy chọn ít nhất 1 câu.")
    db=SessionLocal()
    try:
        selected=[]; shortage=[]
        for kind,levels in matrix.items():
            for level,count in levels.items():
                if not count: continue
                pool=db.query(CauHoi).filter(CauHoi.mon_hoc==mon_hoc.strip(),CauHoi.khoi_lop==khoi_lop,CauHoi.dang_cau_hoi==kind,CauHoi.muc_do==level).all()
                if len(pool)<count: shortage.append(f"{kind} - {level}: cần {count}, có {len(pool)}")
                else: selected.extend(random.sample(pool,count))
        if shortage: raise HTTPException(400,"Ngân hàng chưa đủ câu hỏi:\\n"+"\\n".join(shortage)+"\\nHãy nhập thêm câu hoặc giảm số lượng.")
        questions=[as_dict(q) for q in selected]
    finally: db.close()
    config={"mon_hoc":mon_hoc,"khoi_lop":khoi_lop,"loai_kiem_tra":loai_kiem_tra,"thoi_gian":thoi_gian,"ma_tran":matrix}
    output=Path("/tmp")/f"De_Thi_{uuid.uuid4().hex[:8]}.docx"
    try: tao_file_word_de_thi(config,str(output),questions,kem_dap_an.lower()=="yes")
    except Exception as e: raise HTTPException(500,f"Không xuất được Word: {e}")
    return FileResponse(str(output),filename=f"De_Thi_{mon_hoc.replace(' ','_')}_Khoi_{khoi_lop}.docx",media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
