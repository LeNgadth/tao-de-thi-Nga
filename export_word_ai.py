from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def tao_file_word_de_thi(ma_tran_config, output_filename="De_Thi.docx", generated_questions=None):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = header.add_run(f"ĐỀ KIỂM TRA {ma_tran_config['loai_kiem_tra'].upper()}\\n")
    title.bold = True
    title.font.size = Pt(14)
    header.add_run(
        f"Môn: {ma_tran_config['mon_hoc']} - Khối: {ma_tran_config['khoi_lop']}\\n"
        f"Thời gian làm bài: {ma_tran_config['thoi_gian']} phút"
    )
    doc.add_paragraph("Lưu ý: Giáo viên cần rà soát nội dung và đáp án do AI tạo trước khi sử dụng.")
    questions = generated_questions or []
    stt = 1
    d1 = [q for q in questions if q.get("dang_cau_hoi") == "Dang_I"]
    d2 = [q for q in questions if q.get("dang_cau_hoi") == "Dang_II"]
    d3 = [q for q in questions if q.get("dang_cau_hoi") == "Dang_III"]

    if d1:
        p = doc.add_paragraph()
        p.add_run("PHẦN I. TRẮC NGHIỆM NHIỀU LỰA CHỌN").bold = True
        for q in d1:
            p = doc.add_paragraph()
            p.add_run(f"Câu {stt}: ").bold = True
            p.add_run(q.get("noi_dung", ""))
            doc.add_paragraph(f"A. {q.get('dap_an_a','')}    B. {q.get('dap_an_b','')}")
            doc.add_paragraph(f"C. {q.get('dap_an_c','')}    D. {q.get('dap_an_d','')}")
            stt += 1

    if d2:
        p = doc.add_paragraph()
        p.add_run("PHẦN II. TRẮC NGHIỆM ĐÚNG/SAI").bold = True
        for q in d2:
            p = doc.add_paragraph()
            p.add_run(f"Câu {stt}: ").bold = True
            p.add_run(q.get("noi_dung", ""))
            doc.add_paragraph(
                f"a) {q.get('dap_an_a','')}\\nb) {q.get('dap_an_b','')}\\n"
                f"c) {q.get('dap_an_c','')}\\nd) {q.get('dap_an_d','')}"
            )
            stt += 1

    if d3:
        p = doc.add_paragraph()
        p.add_run("PHẦN III. TỰ LUẬN").bold = True
        for q in d3:
            p = doc.add_paragraph()
            p.add_run(f"Câu {stt}: ").bold = True
            p.add_run(q.get("noi_dung", ""))
            stt += 1

    doc.add_page_break()
    p = doc.add_paragraph()
    p.add_run("ĐÁP ÁN VÀ GỢI Ý GIẢI").bold = True
    stt = 1
    for q in d1:
        doc.add_paragraph(f"Câu {stt}: {q.get('dap_an_dung','')} — {q.get('giai_thich','')}")
        stt += 1
    for q in d2:
        doc.add_paragraph(f"Câu {stt}: {q.get('dap_an_dung','')} — {q.get('giai_thich','')}")
        stt += 1
    for q in d3:
        doc.add_paragraph(
            f"Câu {stt}: Đáp án/hướng dẫn chấm: {q.get('dap_an_a','') or q.get('dap_an_dung','')} "
            f"{q.get('giai_thich','')}"
        )
        stt += 1

    doc.save(output_filename)
    return output_filename
